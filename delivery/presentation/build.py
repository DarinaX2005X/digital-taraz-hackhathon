"""Build Tirlik's editable PPTX and vector PDF from the same authored scene model.

No browser, screenshots, network, AI calls, application process, or database.
Python 3.11+; dependencies in requirements.txt. All geometry uses a 1440 x 810 canvas.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
W, H = 1440, 810
SCALE = 2 / 3  # 960 x 540 points, widescreen 13 1/3 x 7 1/2 inches.
C = dict(bg='F3F6F4', white='FFFFFF', ink='17332E', muted='4D655E', action='086A60',
         pale='DDEFE6', inset='EAF0EC', line='C5D2CA', coral='E48068',
         warm='FFF3D8', warning='785414', shadow='E0E8E2', dark='102B26')
APP_URL = 'https://tirlik.vercel.app'
REPO_URL = 'https://github.com/DarinaX2005X/digital-taraz-hackhathon'

def register_fonts(directory):
    for family, filename in [('Noto', 'NotoSans-Regular.ttf'), ('NotoBold', 'NotoSans-Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(family, str(directory / filename)))


class Scene:
    def __init__(self, title, bg=C['bg'], notes='', evidence=()):
        self.title, self.bg, self.notes, self.evidence = title, bg, notes, list(evidence)
        self.objects = []

    def add(self, kind, **kwargs):
        self.objects.append(dict(kind=kind, **kwargs))

    def rect(self, x, y, w, h, fill, r=0, stroke=None, sw=1):
        self.add('rect', x=x, y=y, w=w, h=h, fill=fill, radius=r, stroke=stroke, sw=sw)

    def text(self, text, x, y, w, size=26, color=C['ink'], bold=False, h=None, align='left', link=None):
        lines = text.split('\n')
        self.add('text', text=text, x=x, y=y, w=w, h=h or (len(lines) * size * 1.25 + 6),
                 size=size, color=color, bold=bold, align=align, leading=1.25, link=link)

    def line(self, x1, y1, x2, y2, color=C['line'], width=2):
        self.add('line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width)

    def poly(self, points, fill, stroke=None, sw=1):
        self.add('poly', points=points, fill=fill, stroke=stroke, sw=sw)

    def circle(self, x, y, d, fill, stroke=None, sw=1):
        self.add('circle', x=x, y=y, w=d, h=d, fill=fill, stroke=stroke, sw=sw)

    def arrow(self, x1, y1, x2, y2, color=C['action'], width=3, head=11):
        self.line(x1, y1, x2, y2, color, width)
        angle = math.atan2(y2-y1, x2-x1)
        self.poly([(x2, y2), (x2-head*math.cos(angle-.5), y2-head*math.sin(angle-.5)),
                   (x2-head*math.cos(angle+.5), y2-head*math.sin(angle+.5))], color)

    def panel(self, x, y, w, h, fill=C['white'], r=24, shadow=False):
        if shadow:
            self.rect(x+3, y+7, w, h, C['shadow'], r)
        self.rect(x, y, w, h, fill, r)

    def title_text(self, text, dark=False):
        self.text(text, 80, 62, 1260, 54, C['white'] if dark else C['ink'], True)

    def caption(self, text, dark=False):
        self.text(text, 80, 755, 1270, 18, C['pale'] if dark else C['muted'])

    def logo(self, x, y, size=60, name=False, dark=False):
        self.add('image', path='assets/tirlik-mark.png', x=x, y=y, w=size, h=size)
        if name:
            self.text('Tirlik', x+size+16, y+2, 240, size*.65, C['white'] if dark else C['ink'], True)


def paper(s, x, y, w, h, title, label=None):
    s.rect(x+8, y+10, w, h, C['shadow'], 18)
    s.rect(x+3, y+4, w, h, C['inset'], 18)
    s.rect(x, y, w, h, C['white'], 18, C['line'], 1)
    s.text(title, x+26, y+26, w-52, 29, bold=True)
    if label:
        s.text(label, x+26, y+h-48, w-52, 20, C['muted'])


def lock(s, x, y, scale=1, color=C['action']):
    s.rect(x+5*scale, y+10*scale, 20*scale, 16*scale, color, 4*scale)
    s.line(x+10*scale, y+12*scale, x+10*scale, y+4*scale, color, 3*scale)
    s.line(x+10*scale, y+4*scale, x+20*scale, y+4*scale, color, 3*scale)
    s.line(x+20*scale, y+4*scale, x+20*scale, y+12*scale, color, 3*scale)


def check(s, x, y, scale=1, color=C['action']):
    s.line(x, y+9*scale, x+7*scale, y+16*scale, color, 3*scale)
    s.line(x+7*scale, y+16*scale, x+22*scale, y, color, 3*scale)


def storefront(s, x, y, scale=1):
    # Original explanatory vector object, not an application screenshot or business photograph.
    p = lambda v: v*scale
    s.poly([(x+p(10), y+p(155)), (x+p(210), y+p(155)), (x+p(238), y+p(172)), (x+p(37), y+p(172))], C['shadow'])
    s.poly([(x+p(188), y+p(26)), (x+p(216), y+p(40)), (x+p(216), y+p(154)), (x+p(188), y+p(144))], C['action'])
    s.rect(x+p(12), y+p(26), p(178), p(130), C['white'], p(8), C['line'], 1)
    s.rect(x+p(28), y+p(66), p(87), p(74), C['pale'], p(3))
    s.rect(x+p(127), y+p(66), p(46), p(90), C['inset'], p(3))
    s.line(x+p(149), y+p(66), x+p(149), y+p(156), C['line'], p(2))
    s.rect(x+p(2), y+p(18), p(197), p(44), C['action'], p(8))
    s.rect(x+p(20), y+p(31), p(95), p(10), C['white'], p(4))
    s.rect(x+p(124), y+p(31), p(54), p(10), C['coral'], p(4))
    s.poly([(x+p(32), y+p(71)), (x+p(69), y+p(71)), (x+p(108), y+p(113)), (x+p(108), y+p(135)), (x+p(92), y+p(135))], 'EDF7F2')


def scenes():
    from product_scenes import build_scenes
    return build_scenes(Scene, W, H)


def validate_scene(slides):
    warnings=[]; texts=0
    for si,s in enumerate(slides,1):
        for oi,o in enumerate(s.objects):
            if o['kind'] in ('rect','circle','image','text'):
                if min(o['x'],o['y']) < -.1 or o['x']+o['w']>W+.1 or o['y']+o['h']>H+.1:
                    warnings.append(f'slide {si} object {oi}: outside slide')
            if o['kind']=='text':
                texts+=1
                f='NotoBold' if o['bold'] else 'Noto'
                for ln in o['text'].split('\n'):
                    width=pdfmetrics.stringWidth(ln,f,o['size'])
                    if width>o['w']-2: warnings.append(f'slide {si}: line width {width:.1f} > {o["w"]}: {ln}')
                    missing=[c for c in ln if not c.isspace() and ord(c) not in pdfmetrics.getFont(f).face.charWidths]
                    if missing: warnings.append(f'slide {si}: missing glyphs {missing}')
                if '\u2014' in o['text'] or '\u2013' in o['text']:
                    warnings.append(f'slide {si}: authored dash punctuation')
    if warnings: raise ValueError('\n'.join(warnings))
    return {'slides':len(slides),'editable_text_boxes':texts,'geometry_and_font_checks':'passed'}


def pdf(slides,path):
    cv=canvas.Canvas(str(path),pagesize=(W*SCALE,H*SCALE),pageCompression=1,invariant=1)
    cv.setTitle('Tirlik — Всё для вашего бизнеса')
    cv.setAuthor('Tirlik')
    cv.setSubject('Рабочие материалы по данным и ограничениям бизнеса')
    for s in slides:
        cv.setFillColor(HexColor('#'+s.bg)); cv.rect(0,0,W*SCALE,H*SCALE,fill=1,stroke=0)
        for o in s.objects:
            k=o['kind']
            cv.saveState()
            if k in ('rect','circle'):
                cv.setFillColor(HexColor('#'+o['fill']))
                cv.setStrokeColor(HexColor('#'+(o['stroke'] or o['fill'])))
                cv.setLineWidth(o['sw']*SCALE)
                args=(o['x']*SCALE,(H-o['y']-o['h'])*SCALE,o['w']*SCALE,o['h']*SCALE)
                if k=='circle': cv.ellipse(args[0],args[1],args[0]+args[2],args[1]+args[3],fill=1,stroke=bool(o['stroke']))
                else: cv.roundRect(*args,o['radius']*SCALE,fill=1,stroke=bool(o['stroke']))
            elif k=='line':
                cv.setStrokeColor(HexColor('#'+o['color']));cv.setLineWidth(o['width']*SCALE);cv.setLineCap(1)
                cv.line(o['x1']*SCALE,(H-o['y1'])*SCALE,o['x2']*SCALE,(H-o['y2'])*SCALE)
            elif k=='poly':
                p=cv.beginPath();p.moveTo(o['points'][0][0]*SCALE,(H-o['points'][0][1])*SCALE)
                for x,y in o['points'][1:]:p.lineTo(x*SCALE,(H-y)*SCALE)
                p.close();cv.setFillColor(HexColor('#'+o['fill']));cv.setStrokeColor(HexColor('#'+(o['stroke'] or o['fill'])));cv.setLineWidth(o['sw']*SCALE)
                cv.drawPath(p,fill=1,stroke=bool(o['stroke']))
            elif k=='image':
                cv.drawImage(str(ROOT/o['path']),o['x']*SCALE,(H-o['y']-o['h'])*SCALE,o['w']*SCALE,o['h']*SCALE,mask='auto')
            elif k=='text':
                font='NotoBold' if o['bold'] else 'Noto';cv.setFont(font,o['size']*SCALE);cv.setFillColor(HexColor('#'+o['color']))
                for i,ln in enumerate(o['text'].split('\n')):
                    y=(H-o['y']-o['size']*.98-i*o['size']*o['leading'])*SCALE
                    x=o['x']*SCALE
                    if o['align']=='center':cv.drawCentredString(x+o['w']*SCALE/2,y,ln)
                    elif o['align']=='right':cv.drawRightString(x+o['w']*SCALE,y,ln)
                    else:cv.drawString(x,y,ln)
                if o['link']:cv.linkURL(o['link'],(o['x']*SCALE,(H-o['y']-o['h'])*SCALE,(o['x']+o['w'])*SCALE,(H-o['y'])*SCALE),relative=0,thickness=0)
            cv.restoreState()
        cv.showPage()
    cv.save()


def pptx(slides,path):
    prs=Presentation();prs.slide_width=Pt(W*SCALE);prs.slide_height=Pt(H*SCALE)
    prs.core_properties.title='Tirlik | Всё для вашего бизнеса';prs.core_properties.author='Tirlik'
    prs.core_properties.subject='Рабочие материалы по данным и ограничениям бизнеса'
    for s in slides:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(s.bg)
        for i,o in enumerate(s.objects):
            k=o['kind'];shape=None
            if k in ('rect','circle'):
                typ=MSO_AUTO_SHAPE_TYPE.OVAL if k=='circle' else MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if o.get('radius') else MSO_AUTO_SHAPE_TYPE.RECTANGLE
                shape=slide.shapes.add_shape(typ,*[Pt(o[v]*SCALE) for v in ('x','y','w','h')])
                if k=='rect' and o.get('radius'):shape.adjustments[0]=min(.5,o['radius']/min(o['w'],o['h']))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(o['fill'])
                if o['stroke']:shape.line.color.rgb=RGBColor.from_string(o['stroke']);shape.line.width=Pt(o['sw']*SCALE)
                else:shape.line.fill.background()
            elif k=='line':
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Pt(o[v]*SCALE) for v in ('x1','y1','x2','y2')])
                shape.line.color.rgb=RGBColor.from_string(o['color']);shape.line.width=Pt(o['width']*SCALE)
            elif k=='poly':
                pts=[(round(x*100),round(y*100)) for x,y in o['points']]
                f=slide.shapes.build_freeform(*pts[0],scale=Pt(SCALE/100));f.add_line_segments(pts[1:],close=True)
                shape=f.convert_to_shape();shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(o['fill'])
                if o['stroke']:shape.line.color.rgb=RGBColor.from_string(o['stroke']);shape.line.width=Pt(o['sw']*SCALE)
                else:shape.line.fill.background()
            elif k=='image':shape=slide.shapes.add_picture(str(ROOT/o['path']),*[Pt(o[v]*SCALE) for v in ('x','y','w','h')])
            elif k=='text':
                shape=slide.shapes.add_textbox(*[Pt(o[v]*SCALE) for v in ('x','y','w','h')]);tf=shape.text_frame
                tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.TOP
                for j,ln in enumerate(o['text'].split('\n')):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(o['size']*o['leading']*SCALE)
                    p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[o['align']]
                    r=p.add_run();r.text=ln;r.font.name='Noto Sans';r.font.size=Pt(o['size']*SCALE);r.font.bold=o['bold'];r.font.color.rgb=RGBColor.from_string(o['color'])
                    if o['link']:r.hyperlink.address=o['link']
            if shape:shape.name=(o.get('text','').replace('\n',' ')[:70] or k)+f' [{i+1}]'
        slide.notes_slide.notes_text_frame.text=s.notes+'\n\nИсточники механизма и данных:\n'+'\n'.join(s.evidence)
    prs.save(path)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--output',type=Path,default=ROOT)
    ap.add_argument('--font-dir',type=Path,default=ROOT.parents[1]/'public'/'fonts'/'noto-sans')
    args=ap.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    register_fonts(args.font_dir)
    slides=scenes(); checks=validate_scene(slides)
    for ext,fn in [('pdf',pdf),('pptx',pptx)]:
        fn(slides,args.output/f'tirlik-presentation.{ext}')
    from svg_export import export_svg
    for i,s in enumerate(slides,1):
        export_svg(s,args.output/'slides'/f'slide-{i:02d}.svg',ROOT,W,H)
    source={'canvas':[W,H],'identity':{'name':'Tirlik','slogan':'Всё для вашего бизнеса','font':'Noto Sans'},'slides':[vars(s) for s in slides]}
    (args.output/'scene-model.json').write_text(json.dumps(source,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    checks['files']={name:{'bytes':(args.output/name).stat().st_size,'sha256':hashlib.sha256((args.output/name).read_bytes()).hexdigest()} for name in ['tirlik-presentation.pptx','tirlik-presentation.pdf']}
    import pypdfium2 as pdfium
    previews=args.output/'previews';previews.mkdir(parents=True,exist_ok=True)
    document=pdfium.PdfDocument(args.output/'tirlik-presentation.pdf')
    thumbs=[]
    for i in range(len(document)):
        page=document[i]
        bitmap=page.render(scale=1.5)
        rendered=bitmap.to_pil()
        rendered.save(previews/f'slide-{i+1:02d}.png')
        thumbs.append(rendered.resize((480,270),Image.Resampling.LANCZOS))
        bitmap.close();page.close()
    document.close()
    overview=Image.new('RGB',(1488,858),'#03100c')
    for i,im in enumerate(thumbs):overview.paste(im,(12+(i%3)*492,12+(i//3)*282))
    overview.save(previews/'overview.png')
    checks['preview_source']='Actual generated PDF rasterized with PDFium; not an application screenshot or PowerPoint viewer export.'
    (args.output/'build-checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(checks,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
